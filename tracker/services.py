"""The transactional boundary for imports, allocation, and retained legacy history.

All phone-history mutations take a PostgreSQL transaction advisory lock. This
small, internal application favors a simple, auditable correctness boundary over
parallel write throughput. SQLite is for single-process development only.
"""

from contextlib import contextmanager
from datetime import date, datetime
from io import BytesIO
import math
import re
import unicodedata
from zipfile import BadZipFile, ZipFile

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import connection, transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl import load_workbook
import phonenumbers

from .models import (
    Assignment, AuditEvent, CallAttempt, ImportBatch, ImportRow, PhoneNumber, RepBatch,
)
from . import texting_lists


_PHONE_WRITE_LOCK = 73621001
# Dash look-alikes that word processors and web pages put into copied numbers.
_DASHES = str.maketrans({character: "-" for character in "\u2010\u2011\u2012\u2013\u2014\u2015\u2212\ufe58\ufe63"})
_HEADERS = {
    "phone", "phones", "phonenumber", "phonenumbers", "telephone", "telephonenumber",
    "mobile", "mobilenumber", "cell", "cellphone", "number", "contactnumber",
}


@contextmanager
def _phone_transaction():
    with transaction.atomic():
        if connection.vendor == "postgresql":
            with connection.cursor() as cursor:
                cursor.execute("SELECT pg_advisory_xact_lock(%s)", [_PHONE_WRITE_LOCK])
        yield


def _require_manager(actor):
    if not (
        actor and actor.is_authenticated and actor.is_active
        and (actor.is_staff or actor.is_superuser)
    ):
        raise PermissionDenied("Only an administrator can manage lead lists.")


def _require_owner(actor, batch):
    if not (
        actor and actor.is_authenticated and actor.is_active and actor.pk == batch.rep_id
    ):
        raise PermissionDenied("This batch is assigned to another representative.")


def write_audit(actor, action, target, details=None):
    """Append metadata, never spreadsheet content, phone numbers, or credentials."""
    if details is not None and not isinstance(details, dict):
        raise ValueError("Audit details must be a dictionary.")
    target_type = target._meta.model_name if hasattr(target, "_meta") else "system"
    target_id = str(target.pk if hasattr(target, "pk") else target)
    return AuditEvent.objects.create(
        actor=actor if actor and getattr(actor, "is_authenticated", False) else None,
        action=action, target_type=target_type, target_id=target_id[:80], details=details or {},
    )


def _raw_text(raw):
    if isinstance(raw, bool) or raw is None:
        return "" if raw is None else str(raw)
    if isinstance(raw, float) and math.isfinite(raw) and raw.is_integer():
        return str(int(raw))
    return str(raw).strip()


def normalize_phone(raw):
    """Normalize complete US/Canadian numbers without guessing missing digits."""
    if isinstance(raw, bool) or isinstance(raw, (date, datetime)):
        raise ValidationError("Enter a complete US or Canadian phone number.")
    if isinstance(raw, float) and (not math.isfinite(raw) or not raw.is_integer()):
        raise ValidationError("Phone numbers cannot contain decimal values.")
    value = _raw_text(raw)
    if not value or len(value) > 120:
        raise ValidationError("Enter a complete US or Canadian phone number.")
    # Copy-paste artifacts: invisible format marks (Outlook/Teams direction marks, zero-width
    # spaces, BOM), full-width characters, fancy dashes, tabs and odd spaces.
    value = "".join(ch for ch in value if unicodedata.category(ch) != "Cf")
    value = re.sub(r"\s+", " ", unicodedata.normalize("NFKC", value).translate(_DASHES)).strip()
    if not re.fullmatch(r"\+?[0-9\s().\-]+", value):
        raise ValidationError("Use digits with an optional +1, spaces, parentheses, or dashes.")
    digits = re.sub(r"\D", "", value)
    if len(digits) not in (10, 11) or (len(digits) == 11 and not digits.startswith("1")):
        raise ValidationError("Include all 10 digits, with an optional country code of 1.")
    if value.startswith("+") and (len(digits) != 11 or not digits.startswith("1")):
        raise ValidationError("Use +1 followed by the complete 10-digit number.")
    try:
        parsed = phonenumbers.parse(value, "US")
    except phonenumbers.NumberParseException as error:
        raise ValidationError("This phone number could not be read.") from error
    if parsed.country_code != 1 or not phonenumbers.is_valid_number(parsed):
        raise ValidationError("This is not a valid US or Canadian phone number.")
    # Shared +1 non-geographic services (such as toll-free) are valid as well.
    if phonenumbers.region_code_for_number(parsed) not in {"US", "CA", "001"}:
        raise ValidationError("Only US and Canadian numbers are supported.")
    return phonenumbers.format_number(parsed, phonenumbers.PhoneNumberFormat.E164)


def _read_workbook(file):
    filename = str(getattr(file, "name", "")).replace("\\", "/").split("/")[-1]
    if not filename.lower().endswith(".xlsx") or len(filename) > 255:
        raise ValidationError("Upload an .xlsx workbook with a single phone-number column.")
    max_bytes = getattr(settings, "IMPORT_MAX_BYTES", 5 * 1024 * 1024)
    if getattr(file, "size", 0) > max_bytes:
        raise ValidationError("The workbook must be 5 MB or smaller.")
    file.seek(0)
    content = file.read(max_bytes + 1)
    if len(content) > max_bytes:
        raise ValidationError("The workbook must be 5 MB or smaller.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            max_expanded = getattr(settings, "IMPORT_MAX_UNCOMPRESSED_BYTES", 32 * 1024 * 1024)
            if len(entries) > 2000 or sum(item.file_size for item in entries) > max_expanded:
                raise ValidationError("The workbook expands beyond the supported size limit.")
            for item in entries:
                if item.flag_bits & 1:
                    raise ValidationError("Password-protected workbooks are not supported.")
                if item.file_size > max(1024 * 1024, item.compress_size * 200):
                    raise ValidationError("The workbook has an unsupported compression ratio.")
                if item.filename.lower().endswith("vbaproject.bin"):
                    raise ValidationError("Macro-enabled workbooks are not supported.")
    except BadZipFile as error:
        raise ValidationError("This file is not a valid .xlsx workbook.") from error

    workbook = None
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        max_rows = getattr(settings, "IMPORT_MAX_ROWS", 10000)
        parsed_rows = []
        populated_sheets = 0
        for sheet in workbook.worksheets:
            # Do not trust a user-controlled worksheet dimension to hide cells.
            sheet.reset_dimensions()
            sheet_has_values = False
            source_column = None
            first_populated_row = True
            for row_number, cells in enumerate(sheet.iter_rows(), start=1):
                # Styled but empty rows don't count toward max_rows (checked on parsed rows
                # below); the uncompressed-size limit above bounds this loop, and no real
                # worksheet has more rows than Excel allows.
                if row_number > 1_048_576:
                    raise ValidationError(f"A workbook can contain at most {max_rows:,} phone rows.")
                if any(cell.data_type == "f" for cell in cells):
                    raise ValidationError("Formulas are not allowed. Paste the phone numbers as values.")
                populated = [(index, cell) for index, cell in enumerate(cells) if _raw_text(cell.value)]
                if not populated:
                    continue
                sheet_has_values = True
                if len(populated) != 1:
                    raise ValidationError("Use a single column of phone numbers, with no other data columns.")
                column, cell = populated[0]
                if source_column is None:
                    source_column = column
                elif source_column != column:
                    raise ValidationError("Keep every phone number in the same column.")
                value = _raw_text(cell.value)
                if first_populated_row:
                    first_populated_row = False
                    # A header is any first row without digits, or a known name like "Phone 1".
                    if not any(character.isdigit() for character in value) or re.sub(r"[^a-z]", "", value.lower()) in _HEADERS:
                        continue
                try:
                    canonical = normalize_phone(cell.value)
                    error = ""
                except ValidationError as exception:
                    canonical = None
                    error = exception.messages[0]
                parsed_rows.append((row_number, value[:120], canonical, error))
                if len(parsed_rows) > max_rows:
                    raise ValidationError(f"A workbook can contain at most {max_rows:,} phone rows.")
            if sheet_has_values:
                populated_sheets += 1
                if populated_sheets > 1:
                    raise ValidationError("Use one worksheet containing the phone-number column.")
        if not parsed_rows:
            raise ValidationError("The workbook has no phone-number rows.")
        return filename, parsed_rows
    except ValidationError:
        raise
    except Exception as error:
        # Parser errors are deliberately not exposed to the browser.
        raise ValidationError("The workbook could not be read. Save it as a new .xlsx file and try again.") from error
    finally:
        if workbook is not None:
            workbook.close()


def _validate_list_type(list_type, *, active_only=True):
    """A texting list key that exists (and, for new work, is not hidden)."""
    keys = texting_lists.active_keys() if active_only else set(texting_lists.all_lists())
    if not isinstance(list_type, str) or list_type not in keys:
        raise ValidationError("Choose one of your texting lists for this upload.")
    return list_type


def eligible_reps(list_type):
    """Current active salesperson accounts belonging to the given texting list."""
    list_type = _validate_list_type(list_type, active_only=False)
    return get_user_model().objects.filter(
        is_active=True, is_staff=False, is_superuser=False,
        list_memberships__list_type=list_type,
    )


def import_workbook(file, actor, label="", *, list_type):
    _require_manager(actor)
    list_type = _validate_list_type(list_type)
    filename, parsed_rows = _read_workbook(file)
    label = str(label).strip() or filename[:-5]
    if len(label) > 120:
        raise ValidationError("The list name must be 120 characters or fewer.")
    with _phone_transaction():
        all_numbers = list(dict.fromkeys(row[2] for row in parsed_rows if row[2]))
        existing = PhoneNumber.objects.in_bulk(all_numbers, field_name="canonical")
        latest_assignments = {}
        existing_ids = [phone.pk for phone in existing.values()]
        for offset in range(0, len(existing_ids), 500):
            assignments = Assignment.objects.filter(
                phone_id__in=existing_ids[offset:offset + 500]
            ).order_by("phone_id", "-batch__created_at", "-pk").values(
                "phone_id", "batch__scheduled_date", "batch__rep__username",
            )
            for assignment in assignments:
                latest_assignments.setdefault(assignment["phone_id"], assignment)
        new_numbers = [PhoneNumber(canonical=number) for number in all_numbers if number not in existing]
        PhoneNumber.objects.bulk_create(new_numbers, batch_size=500)
        phones = PhoneNumber.objects.in_bulk(all_numbers, field_name="canonical")
        upload = ImportBatch.objects.create(
            label=label, filename=filename, uploaded_by=actor, list_type=list_type,
        )
        rows = []
        seen = set()
        counts = {choice: 0 for choice in ImportRow.Classification.values}
        for row_number, raw, canonical, error in parsed_rows:
            phone = phones.get(canonical)
            if canonical is None:
                classification, message = ImportRow.Classification.INVALID, error
            elif canonical in seen:
                classification = ImportRow.Classification.DUPLICATE
                message = "Duplicate in this file; no additional assignment will be created."
            elif canonical in existing:
                previous_assignment = latest_assignments.get(phone.pk)
                first_uploaded = timezone.localtime(phone.first_seen).date().isoformat()
                message = f"Phone number uploaded previously on {first_uploaded}."
                if previous_assignment:
                    rep_name = previous_assignment["batch__rep__username"][:60]
                    assigned_date = previous_assignment["batch__scheduled_date"].isoformat()
                    message += f" Assigned to {rep_name} for {assigned_date}."
                # Retain legacy suppression classifications, without surfacing sales
                # outcomes or notes in the list manager's import report.
                classification = (
                    ImportRow.Classification.BLOCKED if phone.status == PhoneNumber.Status.BLOCKED
                    else ImportRow.Classification.PREVIOUS
                )
            else:
                classification, message = ImportRow.Classification.READY, "New number, ready for allocation."
            if canonical:
                seen.add(canonical)
            counts[classification] += 1
            rows.append(ImportRow(
                upload=upload, row_number=row_number, raw_value=raw, phone=phone,
                classification=classification, message=message,
            ))
        ImportRow.objects.bulk_create(rows, batch_size=500)
        write_audit(actor, "upload.imported", upload, {
            "counts": counts, "total_rows": len(rows), "list_type": list_type,
        })
        return upload


def classify_upload(upload, list_type, actor):
    """Choose a draft upload's list type without replacing its phone history."""
    _require_manager(actor)
    list_type = _validate_list_type(list_type)
    with _phone_transaction():
        current = ImportBatch.objects.select_for_update().get(pk=upload.pk)
        if current.status != ImportBatch.Status.DRAFT or current.batches.exists():
            raise ValidationError("Only an unassigned draft upload can change its texting list type.")
        if current.list_type == list_type:
            return current
        previous_type = current.list_type
        current.list_type = list_type
        current.save(update_fields=["list_type"])
        write_audit(actor, "upload.type_changed", current, {
            "list_type": list_type, "previous_list_type": previous_type,
        })
        return current


def reusable_previous_rows(upload):
    """Previously uploaded rows the manager may choose to assign again.

    Legacy do-not-contact numbers are never offered for reuse.
    """
    return upload.rows.filter(classification=ImportRow.Classification.PREVIOUS).exclude(
        phone__status=PhoneNumber.Status.BLOCKED,
    )


def assignable_rows(upload):
    """Rows a split assigns: new numbers, plus previously uploaded ones if the manager opted in."""
    eligible = Q(classification=ImportRow.Classification.READY, phone__status=PhoneNumber.Status.NEW)
    if upload.include_previous:
        eligible |= Q(classification=ImportRow.Classification.PREVIOUS) & ~Q(phone__status=PhoneNumber.Status.BLOCKED)
    return upload.rows.filter(eligible)


def _require_previous_decision(upload):
    if upload.include_previous is None and reusable_previous_rows(upload).exists():
        raise ValidationError("Choose whether to upload the previously uploaded numbers again before assigning this list.")


def decide_previous_numbers(upload, include, actor):
    """Record whether a draft's previously uploaded numbers join its split."""
    _require_manager(actor)
    if not isinstance(include, bool):
        raise ValidationError("Choose Yes or No.")
    with _phone_transaction():
        current = ImportBatch.objects.select_for_update().get(pk=upload.pk)
        if current.status != ImportBatch.Status.DRAFT or current.batches.exists():
            raise ValidationError("Only an unassigned draft upload can change which numbers it includes.")
        if current.include_previous is include:
            return current
        current.include_previous = include
        current.save(update_fields=["include_previous"])
        write_audit(actor, "upload.previous_numbers_decided", current, {
            "include": include, "previous_count": reusable_previous_rows(current).count(),
        })
        return current


def _clean_split(dates, reps, *, list_type, enforce_future=True):
    # Hiding a texting list stops new uploads; drafts already on it can still be split.
    list_type = _validate_list_type(list_type, active_only=False)
    try:
        clean_dates = [date.fromisoformat(item) if isinstance(item, str) else item for item in dates]
    except (ValueError, TypeError) as error:
        raise ValidationError("Choose valid list dates.") from error
    if not clean_dates or any(type(item) is not date for item in clean_dates):
        raise ValidationError("Choose at least one valid list date.")
    if len(clean_dates) != len(set(clean_dates)):
        raise ValidationError("Each list date can be selected only once.")
    if len(clean_dates) > getattr(settings, "MAX_SPLIT_DAYS", 31):
        raise ValidationError("Choose no more than 31 list dates.")
    clean_dates.sort()
    if enforce_future and clean_dates[0] < timezone.localdate():
        raise ValidationError("Calling dates cannot be in the past.")
    try:
        rep_ids = [item.pk if hasattr(item, "pk") else int(item) for item in reps]
    except (ValueError, TypeError) as error:
        raise ValidationError("Choose active representative accounts.") from error
    if not rep_ids or len(rep_ids) > getattr(settings, "MAX_SPLIT_REPS", 100):
        raise ValidationError("Choose between 1 and 100 representatives.")
    if len(rep_ids) != len(set(rep_ids)):
        raise ValidationError("Each representative can be selected only once.")
    users = eligible_reps(list_type).filter(pk__in=rep_ids).in_bulk()
    if len(users) != len(rep_ids):
        label = texting_lists.label(list_type)
        raise ValidationError(f"Every selected representative must be active and belong to {label}.")
    # Stable account ordering keeps preview, retries, and allocation identical.
    clean_reps = sorted(users.values(), key=lambda rep: (rep.username.casefold(), rep.pk))
    return clean_dates, clean_reps


def _allocation(count, dates, reps):
    """Day-first, then rep, counts; entries with no numbers are left out so nobody gets an empty list."""
    per_day, extra_days = divmod(count, len(dates))
    entries = []
    rotation = 0
    for day_index, scheduled_date in enumerate(dates):
        day_count = per_day + (day_index < extra_days)
        per_rep, extra_reps = divmod(day_count, len(reps))
        extra_indexes = {(rotation + offset) % len(reps) for offset in range(extra_reps)}
        for rep_index, rep in enumerate(reps):
            entries.append({
                "date": scheduled_date, "rep": rep,
                "count": per_rep + (rep_index in extra_indexes),
            })
        rotation = (rotation + extra_reps) % len(reps)
    return [entry for entry in entries if entry["count"]]


def preview_split(upload, dates, reps):
    with _phone_transaction():
        current = ImportBatch.objects.get(pk=upload.pk)
        dates, reps = _clean_split(dates, reps, list_type=current.list_type)
        if current.status != ImportBatch.Status.DRAFT:
            raise ValidationError("Only a draft upload can be split into batches.")
        _require_previous_decision(current)
        count = assignable_rows(current).count()
        if not count:
            raise ValidationError("There are no eligible numbers to assign in this upload.")
        return _allocation(count, dates, reps)


def publish_batches(upload, dates, reps, actor):
    _require_manager(actor)
    with _phone_transaction():
        current = ImportBatch.objects.select_for_update().get(pk=upload.pk)
        dates, reps = _clean_split(
            dates, reps, list_type=current.list_type,
            enforce_future=current.status == ImportBatch.Status.DRAFT,
        )
        if current.status == ImportBatch.Status.CLOSED:
            raise ValidationError("This upload has been closed and cannot be reopened.")
        if current.status == ImportBatch.Status.PUBLISHED:
            batches = list(current.batches.select_related("rep", "upload"))
            existing_grid = {(batch.scheduled_date, batch.rep_id) for batch in batches}
            # Replaying the deterministic split reproduces which lists were skipped as empty.
            assigned = Assignment.objects.filter(batch__upload=current).count()
            requested_grid = {(entry["date"], entry["rep"].pk) for entry in _allocation(assigned, dates, reps)}
            # Uploads published before empty lists were skipped have the full grid.
            full_grid = {(day, rep.pk) for day in dates for rep in reps}
            if existing_grid not in (requested_grid, full_grid):
                raise ValidationError("This upload has already been assigned with different split settings.")
            return batches
        _require_previous_decision(current)
        rows = list(assignable_rows(current).select_related("phone"))
        if not rows:
            raise ValidationError("There are no eligible numbers to assign in this upload.")
        # A number has one live assignment; reusing it retires the earlier one, keeping its history.
        reused_ids = [row.phone_id for row in rows if row.classification == ImportRow.Classification.PREVIOUS]
        for offset in range(0, len(reused_ids), 500):
            Assignment.objects.filter(
                phone_id__in=reused_ids[offset:offset + 500],
                status__in=[Assignment.Status.PENDING, Assignment.Status.IN_PROGRESS],
            ).update(status=Assignment.Status.REASSIGNED)
        allocation = _allocation(len(rows), dates, reps)
        batches = []
        assignments = []
        cursor = 0
        for entry in allocation:
            batch = RepBatch.objects.create(upload=current, rep=entry["rep"], scheduled_date=entry["date"])
            batches.append(batch)
            for sequence, row in enumerate(rows[cursor:cursor + entry["count"]], start=1):
                assignments.append(Assignment(batch=batch, phone=row.phone, sequence=sequence))
            cursor += entry["count"]
        Assignment.objects.bulk_create(assignments, batch_size=500)
        phone_ids = [row.phone_id for row in rows]
        # Batch sizes also avoid SQLite's bind-parameter limits in development.
        for offset in range(0, len(phone_ids), 500):
            PhoneNumber.objects.filter(pk__in=phone_ids[offset:offset + 500]).update(
                status=PhoneNumber.Status.RESERVED, updated_at=timezone.now(),
            )
        current.status = ImportBatch.Status.PUBLISHED
        current.save(update_fields=["status"])
        write_audit(actor, "upload.published", current, {
            "dates": [day.isoformat() for day in dates], "rep_ids": [rep.pk for rep in reps],
            "batch_count": len(batches), "assigned_count": len(assignments),
            "reused_count": len(reused_ids), "list_type": current.list_type,
        })
        return batches


def _close_batch_locked(batch, actor):
    now = timezone.now()
    active = batch.assignments.filter(
        status__in=[Assignment.Status.PENDING, Assignment.Status.IN_PROGRESS],
    )
    phone_ids = list(active.values_list("phone_id", flat=True))
    was_closed = batch.status == RepBatch.Status.CLOSED
    if was_closed and not phone_ids:
        return batch
    active.update(status=Assignment.Status.CANCELLED, completed_at=now)
    for offset in range(0, len(phone_ids), 500):
        PhoneNumber.objects.filter(
            pk__in=phone_ids[offset:offset + 500],
            status__in=[PhoneNumber.Status.NEW, PhoneNumber.Status.RESERVED],
        ).update(status=PhoneNumber.Status.RETIRED, updated_at=now)
    if not was_closed:
        batch.status = RepBatch.Status.CLOSED
        batch.closed_at = now
        batch.closed_by = actor
        batch.save(update_fields=["status", "closed_at", "closed_by"])
    write_audit(actor, "batch.closed", batch, {"cancelled_count": len(phone_ids)})
    return batch


def close_batch(batch, actor):
    _require_manager(actor)
    with _phone_transaction():
        current = RepBatch.objects.select_for_update().select_related("upload", "rep").get(pk=batch.pk)
        return _close_batch_locked(current, actor)


def close_upload(upload, actor):
    _require_manager(actor)
    with _phone_transaction():
        current = ImportBatch.objects.select_for_update().get(pk=upload.pk)
        for batch in current.batches.select_for_update().select_related("rep", "upload"):
            _close_batch_locked(batch, actor)
        # Draft imports also retain their history and can never revive on reimport.
        PhoneNumber.objects.filter(
            import_rows__upload=current, import_rows__classification=ImportRow.Classification.READY,
            status=PhoneNumber.Status.NEW,
        ).update(status=PhoneNumber.Status.RETIRED, updated_at=timezone.now())
        if current.status == ImportBatch.Status.CLOSED:
            return current
        current.status = ImportBatch.Status.CLOSED
        current.closed_at = timezone.now()
        current.closed_by = actor
        current.save(update_fields=["status", "closed_at", "closed_by"])
        write_audit(actor, "upload.closed", current, {"batch_count": current.batches.count()})
        return current


def start_call(assignment, actor):
    """Legacy compatibility API; the list-manager UI does not expose calls."""
    with _phone_transaction():
        current = Assignment.objects.select_for_update().select_related(
            "batch__upload", "batch__rep", "phone"
        ).get(pk=assignment.pk)
        _require_owner(actor, current.batch)
        if current.status == Assignment.Status.IN_PROGRESS:
            return current
        if current.status != Assignment.Status.PENDING:
            raise ValidationError("This assignment is no longer available to start.")
        if current.batch.status != RepBatch.Status.OPEN or current.batch.upload.status != ImportBatch.Status.PUBLISHED:
            raise ValidationError("This calling batch has been closed.")
        if current.batch.scheduled_date != timezone.localdate():
            raise PermissionDenied("Calls can only be started on the batch's assigned date.")
        if current.phone.status != PhoneNumber.Status.RESERVED:
            raise ValidationError("This number is no longer available to call.")
        if Assignment.objects.filter(batch__rep=actor, status=Assignment.Status.IN_PROGRESS).exists():
            raise ValidationError("Record the outcome of your current call before starting another.")
        current.status = Assignment.Status.IN_PROGRESS
        current.started_at = timezone.now()
        current.save(update_fields=["status", "started_at"])
        write_audit(actor, "call.started", current, {"batch_id": str(current.batch_id)})
        return current


def record_outcome(assignment, actor, outcome, note=""):
    """Legacy compatibility API; existing outcomes remain immutable history."""
    if outcome not in CallAttempt.Outcome.values:
        raise ValidationError("Choose a valid call outcome.")
    note = str(note).strip()
    if len(note) > 1000:
        raise ValidationError("Call notes must be 1,000 characters or fewer.")
    with _phone_transaction():
        current = Assignment.objects.select_for_update().select_related(
            "batch__upload", "batch__rep", "phone"
        ).get(pk=assignment.pk)
        _require_owner(actor, current.batch)
        if current.status == Assignment.Status.DONE:
            existing = current.attempts.get()
            if existing.outcome == outcome and existing.note == note:
                return current
            raise ValidationError("An outcome has already been recorded for this call.")
        if current.status != Assignment.Status.IN_PROGRESS:
            raise ValidationError("Start the call before recording its outcome.")
        if current.batch.status != RepBatch.Status.OPEN or current.batch.upload.status != ImportBatch.Status.PUBLISHED:
            raise ValidationError("This assignment belongs to a closed batch.")
        CallAttempt.objects.create(assignment=current, actor=actor, outcome=outcome, note=note)
        blocked = outcome in {CallAttempt.Outcome.NOT_INTERESTED, CallAttempt.Outcome.DO_NOT_CALL}
        current.phone.status = PhoneNumber.Status.BLOCKED if blocked else PhoneNumber.Status.CALLED
        current.phone.suppression_reason = outcome if blocked else ""
        current.phone.save(update_fields=["status", "suppression_reason", "updated_at"])
        current.status = Assignment.Status.DONE
        current.completed_at = timezone.now()
        current.save(update_fields=["status", "completed_at"])
        write_audit(actor, "call.outcome_recorded", current, {
            "outcome": outcome, "has_note": bool(note), "previous_status": Assignment.Status.IN_PROGRESS,
            "status": Assignment.Status.DONE,
        })
        return current
