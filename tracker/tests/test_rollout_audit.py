"""Regressions for the pre-rollout audit: each test reproduces a bug that was found and fixed."""
import re
import tempfile
from datetime import timedelta
from io import BytesIO, StringIO
from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import Client
from django.urls import reverse
from openpyxl import Workbook
from openpyxl.styles import Font

from tracker import texting_lists
from tracker.models import ImportBatch, RepBatch, RepListMembership, ScheduleTemplate, TemplateList, TextingList
from tracker.planning import week_start
from tracker.services import import_workbook, normalize_phone
from tracker.tests.test_real_world import PASSWORD, XLSX, Scenario, fake_numbers, xlsx


def members(key):
    return set(RepListMembership.objects.filter(list_type=key).values_list('rep__username', flat=True))


class ThrottledSignInTests(Scenario):
    def attempt(self, password):
        response = Client().post(reverse('login'), {'username': 'anthony.diaz', 'password': password})
        return response, re.sub(r'name="csrfmiddlewaretoken" value="[^"]+"', '', response.content.decode())

    def test_a_locked_out_sign_in_looks_the_same_for_right_and_wrong_passwords(self):
        for _ in range(5):
            self.attempt('wrong')
        wrong, wrong_page = self.attempt('still-wrong')
        right, right_page = self.attempt(PASSWORD)
        self.assertEqual((wrong.status_code, right.status_code), (429, 429))
        self.assertEqual(wrong_page, right_page)
        self.assertIn('Too many sign-in attempts', right_page)
        self.assertNotIn('correct username', wrong_page)
        self.assertIn('value="anthony.diaz"', right_page)


class TextingListMemberTests(Scenario):
    def test_saving_a_lists_reps_keeps_disabled_reps_on_it(self):
        gfs = TextingList.objects.get(key='gfs')
        self.assign(2, [self.today + timedelta(days=1)], [self.blake], list_type='gfs')
        self.blake.is_active = False
        self.blake.save()
        response = self.manager.post(gfs.get_absolute_url(), {'action': 'members', 'reps': [self.emilio.pk, self.anthony.pk]})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(members('gfs'), {'blake.fiorito', 'emilio.arguello', 'anthony.diaz'})

    def test_unticking_an_active_rep_still_removes_them(self):
        gfs = TextingList.objects.get(key='gfs')
        self.manager.post(gfs.get_absolute_url(), {'action': 'members', 'reps': [self.blake.pk]})
        self.assertEqual(members('gfs'), {'blake.fiorito'})


class HiddenListDraftTests(Scenario):
    def test_a_draft_can_still_be_split_after_its_list_is_hidden(self):
        upload = self.upload_via_ui(xlsx(fake_numbers(4)), label='Before hiding', list_type='ringcentral')
        texting_lists.update_list(self.admin, TextingList.objects.get(key='ringcentral'), name='Ringcentral Texting', nickname='Clean', is_active=False)
        response = self.split_via_ui(upload, [self.today], [self.anthony, self.erik])
        self.assertEqual(response.status_code, 302)
        self.assertEqual(RepBatch.objects.filter(upload=upload).count(), 2)
        self.assertEqual(self.split_via_ui(upload, [self.today], [self.anthony, self.erik]).status_code, 302)  # double submit
        self.assertNotIn('ringcentral', dict(texting_lists.active_choices()))  # still closed to new uploads


class PostedUploadPageLinkTests(Scenario):
    def test_links_on_a_previewed_split_go_back_to_the_upload(self):
        upload = self.upload_via_ui(xlsx(fake_numbers(120)), label='Big list')
        preview = self.split_via_ui(upload, [self.today], [self.anthony], intent='preview')
        self.assertEqual(preview.status_code, 200)
        self.assertContains(preview, 'href="?page=2"')
        self.assertRedirects(self.manager.get(reverse('upload_publish', args=[upload.pk]), {'page': 2}),
                             upload.get_absolute_url() + '?page=2')
        self.assertRedirects(self.manager.get(reverse('upload_classify', args=[upload.pk])), upload.get_absolute_url())
        self.assertRedirects(self.manager.get(reverse('upload_publish', args=[upload.pk]), {'page': 'x', 'intent': 'create'}),
                             upload.get_absolute_url())
        self.assertEqual(self.signed_in(self.anthony).get(reverse('upload_publish', args=[upload.pk])).status_code, 403)


class SpreadsheetEdgeTests(Scenario):
    def test_formatted_empty_rows_do_not_count_toward_the_row_limit(self):
        book = Workbook()
        sheet = book.active
        for row in range(1, 10_101):
            sheet.cell(row=row, column=1).font = Font(bold=True)
        for row, number in enumerate(fake_numbers(50), start=1):
            sheet.cell(row=row, column=1, value=number)
        buffer = BytesIO()
        book.save(buffer)
        upload = import_workbook(SimpleUploadedFile('styled.xlsx', buffer.getvalue(), content_type=XLSX), self.admin, 'Styled', list_type='gfs')
        self.assertEqual(upload.rows.count(), 50)

    def test_the_real_row_limit_still_applies(self):
        with self.settings(IMPORT_MAX_ROWS=10):
            self.assertEqual(self.manager.post(reverse('upload_new'), {
                'label': 'Too many', 'list_type': 'gfs', 'file': xlsx(fake_numbers(11))}).status_code, 200)
        self.assertFalse(ImportBatch.objects.exists())

    def test_invisible_copy_paste_marks_are_ignored(self):
        for pasted in ('\u202a+1 (312) 555-0100\u202c', '\u2066312\u200b-555-0100\u2069', '\ufeff3125550100', '312\u2060555\u200d0100'):
            self.assertEqual(normalize_phone(pasted), '+13125550100', repr(pasted))
        upload = import_workbook(xlsx(['\u202a(646) 555-0101\u202c', '\ufeff646.555.0102']), self.admin, 'Pasted', list_type='gfs')
        self.assertEqual(sorted(upload.rows.values_list('phone__canonical', flat=True)), ['+16465550101', '+16465550102'])


class PlannedUploadTypeTests(Scenario):
    def setUp(self):
        super().setUp()
        template = ScheduleTemplate.objects.create(name='Week', is_active=True, created_by=self.admin)
        self.slot = TemplateList.objects.create(template=template, label='GFS midweek', list_type='gfs', weekdays=[1, 2, 3])
        self.slot.reps.set([self.blake])
        self.week = week_start(self.today) + timedelta(weeks=1)

    def upload_for_slot(self, list_type):
        return self.manager.post(reverse('upload_new'), {'label': 'GFS midweek', 'list_type': list_type, 'file': xlsx(fake_numbers(3)),
                                                         'slot': self.slot.pk, 'week': self.week.isoformat()}, follow=True)

    def test_switching_the_texting_list_on_a_planned_upload_is_called_out(self):
        response = self.upload_for_slot('ringcentral')
        self.assertEqual(ImportBatch.objects.get().template_list, self.slot)  # a deliberate swap still fills the plan
        self.assertContains(response, 'Heads up: the planned “GFS midweek” list is for GFS Texting (Donut), but this file went to Ringcentral Texting (Clean).')

    def test_an_upload_to_the_planned_list_has_no_warning(self):
        response = self.upload_for_slot('gfs')
        self.assertEqual(ImportBatch.objects.get().template_list, self.slot)
        self.assertNotContains(response, 'Heads up')


class ProvisioningTests(Scenario):
    def test_provisioning_works_when_a_roster_rep_is_on_a_hidden_list(self):
        old = texting_lists.create_list(self.admin, name='Old promo', reps=[self.blake])
        texting_lists.update_list(self.admin, old, name='Old promo', nickname='', is_active=False)
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / 'roster.txt'
            call_command('provision_texting_reps', admin='ada.admin', output=str(output), stdout=StringIO())
            self.assertTrue(output.exists())
        self.assertEqual(members(old.key), {'blake.fiorito'})
        self.assertIn('blake.fiorito', members('gfs'))
        self.assertEqual(len(members('gfs') | members('ringcentral')), 16)
